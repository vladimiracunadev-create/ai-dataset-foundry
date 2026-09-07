from __future__ import annotations

from pathlib import Path

from ai_dataset_foundry.models import DocumentRecord, SourceInfo
from ai_dataset_foundry.utils.hashing import sha256_file, stable_id


def load_docx(path: Path) -> list[DocumentRecord]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX support requires: uv sync --extra documents") from exc

    source_hash = sha256_file(path)
    doc = Document(str(path))
    blocks: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            blocks.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            blocks.append(" | ".join(values))
    text = "\n\n".join(blocks)
    return [DocumentRecord(
        id=stable_id("doc", str(path.resolve()), source_hash),
        text=text,
        source=SourceInfo(kind="docx", locator=str(path), title=path.name),
        metadata={"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)},
        source_sha256=source_hash,
    )]
